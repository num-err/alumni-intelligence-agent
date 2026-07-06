import openpyxl
import pytest

from app.services.export_service import export_upload
from app.services.ingestion_service import ingest_upload
from config import get_settings


def _make_workbook(tmp_path, headers, rows, name="roster.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    path = tmp_path / name
    wb.save(path)
    return path


@pytest.fixture(autouse=True)
def _output_dir(tmp_path, monkeypatch):
    get_settings.cache_clear()
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path / "output"))
    get_settings.cache_clear()
    yield
    get_settings.cache_clear()


def test_export_preserves_canonical_and_extra_columns(tmp_path, db_session):
    path = _make_workbook(
        tmp_path,
        ["First Name", "Last Name", "Major"],
        [["Jane", "Doe", "Computer Science"]],
    )
    upload = ingest_upload(
        db_session, original_filename="roster.xlsx", stored_filename="roster.xlsx", stored_path=path
    )

    output_path = export_upload(upload)

    assert output_path.exists()
    wb = openpyxl.load_workbook(output_path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]

    assert "First Name" in headers
    assert "Email" in headers
    assert "Country of Origin" in headers
    assert "Current Employer" in headers
    assert "Major" in headers  # preserved extra column

    row = [cell.value for cell in ws[2]]
    row_dict = dict(zip(headers, row))
    assert row_dict["First Name"] == "Jane"
    assert row_dict["Major"] == "Computer Science"


def test_export_formats_evidence_and_review_required(tmp_path, db_session):
    path = _make_workbook(tmp_path, ["First Name"], [["Jane"]])
    upload = ingest_upload(
        db_session, original_filename="roster.xlsx", stored_filename="roster.xlsx", stored_path=path
    )
    record = upload.records[0]
    record.evidence = ["First match", "Second match"]
    record.review_required = True
    db_session.commit()

    output_path = export_upload(upload)

    wb = openpyxl.load_workbook(output_path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    row = dict(zip(headers, [cell.value for cell in ws[2]]))

    assert row["Evidence"] == "First match; Second match"
    assert row["Review Required"] == "Yes"


def test_export_writes_to_configured_output_directory(tmp_path, db_session):
    path = _make_workbook(tmp_path, ["First Name"], [["Jane"]])
    upload = ingest_upload(
        db_session, original_filename="roster.xlsx", stored_filename="roster.xlsx", stored_path=path
    )

    output_path = export_upload(upload)

    assert output_path.parent == get_settings().output_path
    assert output_path.name.startswith("roster_completed_")
